"""Export search results as a citation list: CSV (spreadsheets) or .docx (Word)."""

import csv
import io
from datetime import date

from docx import Document
from docx.shared import Pt

COLUMNS = ["source", "branch", "docket", "decided", "outcome", "discharge_before", "discharge_after",
           "narrative_reason", "tags", "url", "citation", "collected"]

NOTE = ("Outcome, discharge types and issue tags were extracted automatically and may be incorrect. "
        "The official decision governs. This list is a research aid, not legal advice.")


def to_csv(results):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(COLUMNS)
    for r in results:
        writer.writerow([("; ".join(r[c]) if c == "tags" else r[c]) or "" for c in COLUMNS])
    # The "BOM" at the start tells Excel the file is UTF-8, so the → arrow displays correctly.
    return "﻿" + out.getvalue()


def to_docx(results, search_description):
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)

    doc.add_heading("Discharge Upgrade Decisions: Citation List", level=1)
    doc.add_paragraph(f"Search: {search_description}")
    doc.add_paragraph(f"Exported {date.today():%B %-d, %Y}. {len(results)} decision(s).")
    doc.add_paragraph(NOTE).runs[0].italic = True

    for r in results:
        doc.add_paragraph(r["citation"], style="List Number")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
